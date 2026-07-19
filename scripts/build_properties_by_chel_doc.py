from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "brand-strategy" / "properties-by-chel-content-strategy.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(25, 35, 45)
MUTED = RGBColor(90, 98, 108)
LIGHT_FILL = "F2F4F7"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.text = "Properties by Chel Content Strategy"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("Properties by Chel")
    run.font.name = "Calibri"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("Content Strategy and Website Direction")
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.color.rgb = MUTED

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.1), Inches(2.1), Inches(2.1)]
    values = [
        ("Brand Role", "Personal authority and relationship brand"),
        ("Main CTA", "Book a consultation or visit Allabode"),
        ("Primary Audience", "Owners, buyers, tenants, sellers, investors"),
    ]
    for i, (label, value) in enumerate(values):
        cell = table.cell(0, i)
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, LIGHT_FILL)
        set_cell_margins(cell)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        r = cell.paragraphs[0].add_run(label)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = DARK_BLUE
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        r.font.size = Pt(10)
        r.font.color.rgb = INK

    doc.add_paragraph()


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = DARK_BLUE


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_section(doc, title, paragraphs=None, bullets=None, quote=None):
    doc.add_heading(title, level=1)
    for para in paragraphs or []:
        doc.add_paragraph(para)
    if quote:
        add_quote(doc, quote)
    if bullets:
        add_bullets(doc, bullets)


def add_content_pillars(doc):
    doc.add_heading("Social Media Content Pillars", level=1)
    doc.add_paragraph(
        "Properties by Chel should be the social and educational voice of the business. "
        "Its content should make Chel visible, helpful, and trustworthy before someone becomes a client."
    )
    pillars = [
        ("Property Education", "Simple tips for buyers, sellers, landlords, tenants, and investors."),
        ("Appraisal Insights", "Plain-language explanations of value, pricing, comparables, location, and property condition."),
        ("Owner And Landlord Tips", "Practical guidance for occupancy, leasing, tenant screening, and property readiness."),
        ("Market Commentary", "Chel's practical point of view on rental demand, buyer behavior, and investment risk."),
        ("Behind The Brand", "Trust-building stories from site visits, client lessons, fieldwork, and daily real estate practice."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.9)
    table.columns[1].width = Inches(4.4)
    hdr = table.rows[0].cells
    hdr[0].text = "Pillar"
    hdr[1].text = "Role"
    for cell in hdr:
        set_cell_fill(cell, LIGHT_FILL)
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.runs[0].bold = True
            p.runs[0].font.color.rgb = DARK_BLUE
    for pillar, role in pillars:
        row = table.add_row().cells
        row[0].text = pillar
        row[1].text = role
        for cell in row:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_30_day_plan(doc):
    doc.add_heading("30-Day Content Starter Plan", level=1)
    plan = [
        ("Week 1", "Introduce Chel, explain broker vs appraiser, and share a leasing checklist for property owners."),
        ("Week 2", "Post a pricing-before-selling guide, explain value factors, and discuss property management timing."),
        ("Week 3", "Share tenant screening advice, a buyer overpayment warning, and behind-the-scenes fieldwork."),
        ("Week 4", "Publish an OFW property management guide, share a client-safe story, and invite consultations."),
    ]
    for timing, focus in plan:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        label = p.add_run(f"{timing}: ")
        label.bold = True
        label.font.color.rgb = DARK_BLUE
        p.add_run(focus)


def build():
    doc = Document()
    style_document(doc)
    add_title(doc)

    add_section(
        doc,
        "Purpose",
        [
            "Properties by Chel should function as Chel's personal authority and relationship brand. "
            "It should help people understand who Chel is, why they should trust her with real estate decisions, and how they can work with her.",
            "The brand should feel warmer and more personal than Allabode Realty and Appraisal Services. "
            "It should build trust, educate the market, and guide serious inquiries toward Allabode."
        ],
    )

    add_section(
        doc,
        "Recommended Brand Role",
        [
            "Properties by Chel should be positioned as the personal real estate advisory brand of Chel, a licensed Real Estate Broker and Appraiser in the Philippines."
        ],
        quote="Practical real estate guidance from a licensed broker and appraiser, helping buyers, sellers, property owners, tenants, and investors make clearer property decisions in the Philippines.",
    )

    add_section(
        doc,
        "Website Direction",
        [
            "The website should be simple, personal, and credibility-led. It should not try to become a large property marketplace. "
            "Its job is to make people trust Chel and then move qualified leads into the right service path."
        ],
        bullets=["Home", "About Chel", "Work With Me", "Property Guides", "Testimonials or Client Stories", "Contact or Book a Consultation"],
    )

    add_section(
        doc,
        "Home Page Content",
        [
            "The home page should quickly explain Chel's expertise and give visitors clear next steps."
        ],
        quote="Real estate guidance from a licensed broker and appraiser.",
        bullets=[
            "Short personal introduction",
            "Professional credentials",
            "Who Chel helps",
            "Featured property tips or guides",
            "Book a Consultation, Visit Allabode, and Read Property Guides calls to action",
        ],
    )

    add_section(
        doc,
        "About Chel",
        [
            "This should be one of the most important pages on the site. Real estate is trust-driven, and this page should help people feel the person behind the professional service."
        ],
        bullets=[
            "Chel's story and reason for entering real estate",
            "Licensed Real Estate Broker and Appraiser credentials",
            "Experience and service focus",
            "Personal approach to clients",
            "Values such as transparency, practical advice, professionalism, and care",
            "Connection to Allabode as the formal service company",
        ],
    )

    add_section(
        doc,
        "Work With Me",
        [
            "This page should present Chel's advisory role without duplicating the entire Allabode service website."
        ],
        bullets=[
            "Buying guidance",
            "Selling guidance",
            "Leasing guidance",
            "Property owner consultation",
            "Appraisal guidance",
            "Investment property consultation",
        ],
        quote="If you need hands-on brokerage, leasing, property management, or formal appraisal support, my team handles this through Allabode Realty and Appraisal Services.",
    )

    add_section(
        doc,
        "Property Guides",
        [
            "The Property Guides section should become the long-term content engine for Properties by Chel. It should help educate the market while building search visibility and personal authority."
        ],
        bullets=[
            "How to price your property before selling",
            "Leasing checklist for condo owners",
            "Short-term vs long-term leasing in the Philippines",
            "What property owners should know before hiring a property manager",
            "What is a real estate appraisal?",
            "Broker vs appraiser: what is the difference?",
            "How OFWs can manage property in the Philippines",
            "Documents needed when selling property",
            "What affects property value in the Philippines?",
        ],
    )

    add_content_pillars(doc)

    add_section(
        doc,
        "Lead Flow",
        [
            "Properties by Chel should guide visitors into the right next step instead of trying to serve every need on the personal website."
        ],
        bullets=[
            "Personal advice: book a consultation with Chel",
            "Leasing, selling, property management, or appraisal service: go to Allabode",
            "Appraisal technology or data product interest: join the future PropIntel waitlist",
        ],
        quote="Properties by Chel is the personal real estate advisory brand of Chel, founder of Allabode Realty and Appraisal Services.",
    )

    add_30_day_plan(doc)

    add_section(
        doc,
        "Final Recommendation",
        [
            "Keep Properties by Chel active, but make its role clear. It should be the personal trust and education layer. "
            "Allabode should be the company that delivers full-service real estate, leasing, property management, brokerage, and appraisal. "
            "PropIntel should remain separate as the future appraisal technology product."
        ],
        bullets=[
            "Properties by Chel builds trust.",
            "Allabode converts and serves clients.",
            "PropIntel scales appraisal intelligence later.",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
